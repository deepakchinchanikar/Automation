import re
from openpyxl import Workbook


# Function to search for next word after a given target word
def find_next_word(text, target_word):
    # Create a pattern to find the target word followed by a word
    pattern = r'\b' + re.escape(target_word) + r'\b\s+(\w+)'
    matches = re.findall(pattern, text)
    return matches


# Function to process the text file and write to Excel
def process_text_to_excel(text_file, target_word1, target_word2, output_excel):
    # Read the content of the text file
    with open(text_file, 'r') as file:
        text = file.read()

    # Find the next words after the target words
    next_word1 = find_next_word(text, target_word1)
    next_word2 = find_next_word(text, target_word2)
    print(next_word1)
    print(next_word2)

    # Create a new workbook and set the active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Search Results"

    # Insert headers in Excel
    ws['A1'] = f"Next word after '{target_word1}'"
    ws['B1'] = f"Next word after '{target_word2}'"

    # Insert the words in the Excel sheet
    for idx, (word1, word2) in enumerate(zip(next_word1, next_word2), start=2):
        ws[f'A{idx}'] = word1
        ws[f'B{idx}'] = word2

    # Save the workbook to the given output path
    wb.save(output_excel)


# Example usage
text_file = 'input.txt'  # Path to your input text file
target_word1 = 'Rename view'  # The first target word
target_word2 = 'from '  # The second target word
output_excel = 'output.xlsx'  # Path to save the output Excel file

# Call the function to process and generate the Excel file
process_text_to_excel(text_file, target_word1, target_word2, output_excel)
